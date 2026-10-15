"""Build the shareable Phase 1 Word references from their Markdown sources.

Run from any directory using the backend virtual environment, which already
contains python-docx. The output contains fictional screenshots and no secrets.
"""

from __future__ import annotations

import re
import os
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "reference"
VIOLET = "6651C7"
INK = "28243C"
MUTED = "625C76"
PALE = "F0ECFB"
REPO = os.environ.get("DOCUMENTATION_REPOSITORY_URL", "https://github.com/SriTharoon05/UB_HRProcess_Agent/blob/main/").rstrip("/") + "/"
repository_url = urlsplit(REPO)
if (repository_url.scheme != "https" or repository_url.hostname != "github.com"
        or repository_url.username or repository_url.password or repository_url.query or repository_url.fragment):
    raise ValueError("DOCUMENTATION_REPOSITORY_URL must be a GitHub HTTPS source URL without credentials")


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    props.append(item)


def hyperlink(paragraph, label, target):
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), VIOLET)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    node.append(run)
    paragraph._p.append(node)


def inline(paragraph, text, source):
    # Preserve emphasis/code and make reference links useful outside GitHub.
    pattern = r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            label, target = link.groups()
            if not target.startswith(("https://", "http://", "#")):
                path, _, anchor = target.partition("#")
                resolved = (source.parent / path).resolve().relative_to(ROOT)
                target = REPO + resolved.as_posix() + ("#" + anchor if anchor else "")
            hyperlink(paragraph, label, target)
        elif part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def set_metadata(document, kind):
    document.core_properties.title = f"TalentMatch AI — {kind} Reference"
    document.core_properties.subject = "Phase 1 recruitment portal deployment and business handoff"
    document.core_properties.author = "TalentMatch AI Project"
    document.core_properties.keywords = "Phase 1, D365, recruitment, Docker, Supabase, human review"
    document.core_properties.comments = (
        "Generated from the repository Markdown reference. Screenshots use fictional demo data."
    )


def prepare_document(kind):
    document = Document()
    set_metadata(document, kind)
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for level, size in ((1, 17), (2, 13), (3, 11)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(VIOLET)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    document.styles["Title"].font.color.rgb = RGBColor.from_string(INK)
    document.styles["Title"].font.size = Pt(34)
    document.styles["Subtitle"].font.color.rgb = RGBColor.from_string(VIOLET)
    document.styles["Caption"].font.size = Pt(9)
    document.styles["Caption"].font.color.rgb = RGBColor.from_string(MUTED)
    document.styles["Caption"].paragraph_format.space_after = Pt(10)

    header = section.header.paragraphs[0]
    header.text = "TALENTMATCH AI  /  PHASE 1"
    header.style = document.styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"{kind} reference  •  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return document


def cover(document, kind, headings):
    document.add_paragraph("PHASE 1  •  RECRUITMENT INTAKE & EVIDENCE REVIEW", "Subtitle")
    document.add_paragraph("TalentMatch AI", "Title")
    subtitle = document.add_paragraph(f"{kind} Reference")
    subtitle.runs[0].font.size = Pt(22)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(VIOLET)
    document.add_paragraph("Dynamics 365 Finance & Operations + React + FastAPI")
    document.add_paragraph("Deployment handoff • 15 October 2026")
    document.add_paragraph(
        "Current delivery: job publication, LinkedIn candidate intake, private resumes, "
        "AI-assisted evidence comparison and HR review/retry."
    )
    note = document.add_paragraph()
    note.add_run("Scope boundary. ").bold = True
    note.add_run(
        "This is Phase 1. Screening, interviews, offers, onboarding and employee maintenance "
        "are roadmap items. AI supports human review and does not make employment decisions."
    )
    box = document.add_table(rows=2, cols=2)
    box.style = "Table Grid"
    box.cell(0, 0).text = "LOCAL PORTAL"
    box.cell(0, 1).text = "https://localhost:7242"
    box.cell(1, 0).text = "LINKEDIN CALLBACK"
    box.cell(1, 1).text = "https://localhost:7242/auth/linkedin/callback"
    for row in box.rows:
        shade(row.cells[0], PALE)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    document.add_paragraph(
        "Shareable project reference. Contains no production credentials or candidate data. "
        "UI illustrations show fictional demo identities and sample evaluations.",
        "Caption",
    )
    document.add_page_break()
    document.add_heading("Contents", 1)
    for heading in headings:
        document.add_paragraph(heading, "List Bullet")
    document.add_paragraph(
        "Word's Navigation Pane also follows the numbered section headings. "
        "Repository links remain clickable in this document.",
        "Caption",
    )
    document.add_page_break()


def table(document, lines, source):
    records = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    records = [
        row for row in records
        if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in row)
    ]
    result = document.add_table(rows=0, cols=len(records[0]))
    result.style = "Table Grid"
    for index, record in enumerate(records):
        cells = result.add_row().cells
        for cell, text in zip(cells, record):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            inline(cell.paragraphs[0], text, source)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5 if len(record) >= 3 else 9)
                    if index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
            if index == 0:
                shade(cell, VIOLET)
            elif index % 2:
                shade(cell, "F8F6FC")
        properties = result.rows[index]._tr.get_or_add_trPr()
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        properties.append(OxmlElement("w:cantSplit"))
    document.add_paragraph()


def code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1
    props = paragraph._p.get_or_add_pPr()
    background = OxmlElement("w:shd")
    background.set(qn("w:fill"), PALE)
    props.append(background)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


def illustration(document, filename, caption):
    path = ROOT / "docs" / "screenshots" / filename
    if not path.is_file():
        return
    document.add_picture(str(path), width=Inches(6.35))
    document.add_paragraph(caption, "Caption")


def convert(source, kind):
    lines = source.read_text(encoding="utf-8").splitlines()
    headings = [line[3:] for line in lines if line.startswith("## ")]
    document = prepare_document(kind)
    cover(document, kind, headings)
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith(("# ", "Document version:")):
            index += 1
            continue
        if line.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            code(document, "\n".join(block))
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            table(document, block, source)
            continue
        heading = re.match(r"^(#{2,4}) (.+)$", line)
        if heading:
            text = heading.group(2)
            document.add_heading(text, len(heading.group(1)) - 1)
            index += 1
            continue
        bullet = re.match(r"^(?:- |\d+\. )(.+)$", line)
        if bullet:
            style = "List Bullet" if line.startswith("- ") else "List Number"
            inline(document.add_paragraph(style=style), bullet.group(1), source)
        else:
            inline(document.add_paragraph(), line, source)
        index += 1
    document.add_page_break()
    document.add_heading("Illustrations — fictional demo only", 1)
    document.add_paragraph(
        "These browser captures illustrate the Phase 1 interface. Names, contacts, jobs "
        "and evaluation values shown here are fictional demo data, not live provider results."
    )
    illustration(
        document,
        "phase1-candidate-application.png",
        "Figure 1. Candidate confirmation fields, private resume upload and consent for the selected job.",
    )
    document.add_page_break()
    illustration(
        document,
        "phase1-hr-evidence.png",
        "Figure 2. HR-only sample evidence, contextual score and original-resume access. "
        "This is an explicitly labeled fictional evaluation.",
    )
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / f"TalentMatch_AI_{kind}_Reference.docx"
    document.save(target)
    # Reopen the OOXML package to verify basic document and illustration integrity.
    with ZipFile(target) as package:
        assert package.testzip() is None, "Corrupt Word package"
        xml = package.read("word/document.xml").decode("utf-8")
        assert "localhost:7242" in xml and "Phase 1" in xml
        assert sum(name.startswith("word/media/") for name in package.namelist()) == 2
    reopened = Document(target)
    assert len(reopened.paragraphs) > 50 and len(reopened.tables) > 5
    print(f"Generated {target.relative_to(ROOT)} ({target.stat().st_size:,} bytes)")


def main():
    for name, kind in (("technical-reference.md", "Technical"), ("business-reference.md", "Business")):
        convert(ROOT / "docs" / name, kind)


if __name__ == "__main__":
    main()
