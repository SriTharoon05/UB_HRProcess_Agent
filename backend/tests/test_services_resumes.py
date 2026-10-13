import io
import os
import zipfile
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import parse_xml

from app.config import Settings
from app.errors import ConfigurationError, ResumeValidationError
from app.services.resumes import MIME_TYPES, ResumeService, sanitize_filename


@pytest.fixture
def settings(tmp_path):
    return Settings(_env_file=None, demo_mode=True, local_resume_dir=tmp_path / "private-resumes")


def pdf_bytes(text="Projects\nPortal\nBuilt Python APIs."):
    with fitz.open() as doc:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
        return doc.tobytes()


def docx_bytes():
    doc = Document()
    doc.add_heading("Projects", 1)
    doc.add_paragraph("Portal: Built React interfaces.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Employment: Built Python APIs."
    doc.sections[0].header.paragraphs[0].text = "Technical resume"
    doc.sections[0].footer.paragraphs[0].text = "Certifications: Cloud fundamentals"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_real_pdf_and_docx_full_extraction(settings):
    service = ResumeService(settings)
    assert "Built Python APIs." in service.extract(pdf_bytes(), "resume.pdf")
    extracted = service.extract(docx_bytes(), "resume.docx")
    assert "Projects" in extracted
    assert "Built React interfaces." in extracted
    assert "Employment: Built Python APIs." in extracted
    assert "Technical resume" in extracted
    assert "Certifications: Cloud fundamentals" in extracted


def test_docx_text_boxes_and_content_controls_are_not_silently_lost(settings):
    document = Document()
    paragraph = document.add_paragraph("Main experience")
    paragraph._p.append(parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing><w:txbxContent><w:p><w:r><w:t>Text box: Built FastAPI services.</w:t></w:r></w:p></w:txbxContent></w:drawing></w:r>'))
    document._element.body.append(parse_xml('<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sdtContent><w:p><w:r><w:t>Content control: Azure certification.</w:t></w:r></w:p></w:sdtContent></w:sdt>'))
    output = io.BytesIO()
    document.save(output)
    text = ResumeService(settings).extract(output.getvalue(), "formatted.docx")
    assert text.count("Text box: Built FastAPI services.") == 1
    assert "Content control: Azure certification." in text
    assert "Main experience" in text


@pytest.mark.parametrize("data,name,mime", [(b"fake", "resume.pdf", "application/pdf"), (b"%PDF-fake", "resume.docx", MIME_TYPES[".docx"]), (b"test", "resume.exe", "application/octet-stream"), (b"", "resume.pdf", "application/pdf")])
def test_invalid_extensions_magic_and_empty_files(settings, data, name, mime):
    with pytest.raises(ResumeValidationError):
        ResumeService(settings).validate_upload(data, name, mime)


def test_mime_size_scanned_and_corrupt_documents(settings):
    service = ResumeService(settings)
    with pytest.raises(ResumeValidationError, match="MIME"):
        service.validate_upload(pdf_bytes(), "resume.pdf", MIME_TYPES[".docx"])
    settings.max_resume_size_mb = 1
    with pytest.raises(ResumeValidationError, match="smaller"):
        service.validate_upload(b"%PDF-" + b"x" * 1024 * 1024, "resume.pdf", "application/pdf")
    with pytest.raises(ResumeValidationError, match="OCR"):
        service.extract(pdf_bytes(""), "scanned.pdf")
    with pytest.raises(ResumeValidationError, match="corrupted"):
        service.validate_upload(b"%PDF-1.4\nno actual document", "bad.pdf", "application/pdf")


def test_docx_zip_bomb_rejected_before_parse(settings):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("word/document.xml", "x" * 2 * 1024 * 1024)
    with pytest.raises(ResumeValidationError, match="compression ratio"):
        ResumeService(settings).validate_upload(output.getvalue(), "resume.docx", MIME_TYPES[".docx"])


async def test_private_local_storage_opaque_refs_permissions_and_cleanup(settings):
    service = ResumeService(settings)
    data = pdf_bytes()
    reference = await service.store(data, "../../candidate name.pdf")
    assert reference.startswith("local:")
    assert "candidate" not in reference
    assert await service.read(reference) == data
    files = list(settings.local_resume_dir.iterdir())
    assert len(files) == 1
    assert os.stat(files[0]).st_mode & 0o777 == 0o600
    assert os.stat(settings.local_resume_dir).st_mode & 0o777 == 0o700
    with pytest.raises(ResumeValidationError, match="reference"):
        await service.read("local:../../secrets.pdf")
    await service.delete(reference)
    assert not list(settings.local_resume_dir.iterdir())


async def test_production_never_falls_back_to_local(settings):
    settings.demo_mode = False
    with pytest.raises(ConfigurationError, match="Azure Blob"):
        await ResumeService(settings).store(pdf_bytes(), "resume.pdf")
    with pytest.raises(ConfigurationError, match="demo mode"):
        await ResumeService(settings).read("local:" + "a" * 32 + ".pdf")


async def test_azure_requires_private_container_and_uses_actual_sdk_calls(settings):
    data = pdf_bytes()
    blob = SimpleNamespace(upload_blob=AsyncMock(), get_blob_properties=AsyncMock(return_value=SimpleNamespace(size=len(data))), download_blob=AsyncMock(return_value=SimpleNamespace(readall=AsyncMock(return_value=data))), delete_blob=AsyncMock())
    container = SimpleNamespace(get_container_properties=AsyncMock(return_value={"public_access": None}), get_blob_client=Mock(return_value=blob))
    sdk = SimpleNamespace(get_container_client=Mock(return_value=container))
    settings.demo_mode = False
    settings.azure_storage_connection_string = "mocked-sdk-only"
    service = ResumeService(settings, blob_client=sdk)
    ref = await service.store(pdf_bytes(), "resume.pdf")
    assert ref.startswith("azure:")
    assert blob.upload_blob.call_args.kwargs["overwrite"] is False
    assert blob.upload_blob.call_args.kwargs["content_settings"].content_type == "application/pdf"
    assert await service.read(ref) == data
    await service.delete(ref)
    blob.delete_blob.assert_awaited_once()
    container.get_container_properties.return_value = {"public_access": "blob"}
    with pytest.raises(ConfigurationError, match="public access"):
        await service.store(pdf_bytes(), "resume.pdf")


def test_filename_sanitization():
    assert sanitize_filename("..\\folder\\candidate.pdf") == "candidate.pdf"
    assert sanitize_filename("folder/<candidate>.pdf") == "_candidate_.pdf"
